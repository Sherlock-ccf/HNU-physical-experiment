"""Fit a trigonometric model to LN-crystal DC voltage and optical power.

Run in the same directory as data.xlsx:
    python plot_electro_optic.py
"""

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from openpyxl import load_workbook
from scipy.optimize import curve_fit


BASE = Path(__file__).resolve().parent
INPUT = BASE / "data.xlsx"
OUTPUT = BASE / "电压-输出功率曲线.png"


def read_measurements(path: Path) -> tuple[list[float], list[float]]:
    sheet = load_workbook(path, data_only=True).active
    voltage, power = [], []
    for row in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
        if row[0] is None and row[1] is None:
            continue
        if not all(isinstance(value, (int, float)) for value in row):
            raise ValueError(f"非数值测量行：{row}")
        voltage.append(float(row[0]))
        power.append(float(row[1]))
    if len(voltage) < 3 or any(b <= a for a, b in zip(voltage, voltage[1:])):
        raise ValueError("电压数据缺失或未严格递增")
    return voltage, power


def cosine_model(voltage, baseline, amplitude, peak_voltage, half_wave_voltage):
    """Power modulation between adjacent bright and dark states."""
    return baseline + amplitude * np.cos(
        np.pi * (voltage - peak_voltage) / half_wave_voltage
    )


def main() -> None:
    voltage, power = read_measurements(INPUT)
    v_data = np.asarray(voltage)
    p_data = np.asarray(power)
    peak_index = max(range(len(power)), key=power.__getitem__)
    # Choose the measured minimum following the maximum: adjacent extrema on
    # the same voltage scan give the half-wave voltage in this experiment.
    trough_index = min(
        range(peak_index + 1, len(power)), key=power.__getitem__
    )
    v_peak, p_peak = voltage[peak_index], power[peak_index]
    v_trough, p_trough = voltage[trough_index], power[trough_index]
    measured_v_pi = v_trough - v_peak

    parameters, _ = curve_fit(
        cosine_model,
        v_data,
        p_data,
        p0=[(p_peak + p_trough) / 2, (p_peak - p_trough) / 2,
            v_peak, measured_v_pi],
        bounds=([0, 0, 400, 500], [250, 250, 900, 1100]),
    )
    baseline, amplitude, fitted_peak, fitted_v_pi = parameters
    fitted_trough = fitted_peak + fitted_v_pi
    fitted_power = cosine_model(v_data, *parameters)
    r_squared = 1 - np.sum((p_data - fitted_power) ** 2) / np.sum(
        (p_data - p_data.mean()) ** 2
    )

    system_font = Path("C:/Windows/Fonts/simhei.ttf")
    if system_font.exists():
        font_manager.fontManager.addfont(str(system_font))
        plt.rcParams["font.family"] = font_manager.FontProperties(
            fname=str(system_font)
        ).get_name()
    plt.rcParams["axes.unicode_minus"] = False

    fig, ax = plt.subplots(figsize=(10, 5.8), dpi=160)
    ax.scatter(v_data, p_data, color="#2166a3", s=24,
               label="实测数据", zorder=3)
    smooth_voltage = np.linspace(min(voltage), max(voltage), 1000)
    ax.plot(smooth_voltage, cosine_model(smooth_voltage, *parameters),
            color="#c44e35", linewidth=2.1,
            label=f"三角函数拟合（$R^2$ = {r_squared:.4f}）")
    ax.axvline(fitted_peak, color="#c44e35", linestyle="--", alpha=0.55)
    ax.axvline(fitted_trough, color="#c44e35", linestyle="--", alpha=0.55)
    ax.annotate("", xy=(fitted_peak, 242), xytext=(fitted_trough, 242),
                arrowprops={"arrowstyle": "<->", "color": "#45624a"})
    ax.text((fitted_peak + fitted_trough) / 2, 246,
            f"拟合 Vπ = {fitted_v_pi:.0f} V", ha="center", color="#45624a")
    ax.text(0.015, 0.97,
            f"实测极值：{v_peak:g} V → {v_trough:g} V，差值 {measured_v_pi:g} V",
            transform=ax.transAxes, va="top", color="#4b4b4b")
    ax.set(xlabel="晶体驱动电压 / V", ylabel="输出光功率 / μW",
           title="LN 晶体电压与输出光功率：三角函数拟合",
           xlim=(-25, 1740), ylim=(0, 270))
    ax.grid(alpha=0.25)
    ax.legend(loc="upper left", bbox_to_anchor=(0.015, 0.86), frameon=False)
    fig.tight_layout()
    fig.savefig(OUTPUT, dpi=200)
    plt.close(fig)
    print(
        f"已生成：{OUTPUT.name}; 实测极值差 = {measured_v_pi:g} V; "
        f"拟合 Vπ = {fitted_v_pi:.1f} V; R² = {r_squared:.4f}; "
        f"拟合 C = {baseline:.2f} μW, A = {amplitude:.2f} μW, "
        f"Vmax = {fitted_peak:.1f} V"
    )


if __name__ == "__main__":
    main()
