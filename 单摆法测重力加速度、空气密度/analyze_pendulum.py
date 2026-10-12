"""由 data.xlsx 生成单摆实验数据分析和表 2、3、4 的拟合图。"""

from pathlib import Path
import math

import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SHEET = load_workbook(ROOT / "data.xlsx", data_only=True).active
plt.rcParams.update({
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "DejaVu Sans"],
    "axes.unicode_minus": False,
    "figure.dpi": 140,
    "savefig.dpi": 200,
    "font.size": 10,
})


def values(rows, cols):
    return np.array(
        [[SHEET.cell(r, c).value for c in cols] for r in rows], dtype=float
    )


def fit_line(x, y):
    (k, b), cov = np.polyfit(x, y, 1, cov=True)
    predicted = k * x + b
    r2 = 1 - np.sum((y - predicted) ** 2) / np.sum((y - y.mean()) ** 2)
    return k, b, r2, cov


def figure(x, y, k, b, xlabel, ylabel, title, filename, theory=None, note=None):
    fig, ax = plt.subplots(figsize=(7.2, 4.8), layout="constrained")
    ax.scatter(x, y, s=42, color="#1769aa", label="实测数据", zorder=3)
    xx = np.linspace(x.min(), x.max(), 200)
    ax.plot(xx, k * xx + b, color="#d35400", lw=2, label="线性拟合")
    if theory is not None:
        kt, bt = theory
        ax.plot(xx, kt * xx + bt, "--", color="#5b5b5b", lw=1.5, label="理论关系")
    ax.set(xlabel=xlabel, ylabel=ylabel, title=title)
    ax.grid(alpha=0.25)
    ax.legend(loc="best")
    if note:
        ax.text(0.02, 0.97, note, transform=ax.transAxes, va="top",
                bbox={"facecolor": "white", "edgecolor": "#dddddd", "alpha": 0.88})
    fig.savefig(ROOT / filename)
    plt.close(fig)


def main():
    pi = math.pi
    g_ref = 9.791  # 报告模板给出的长沙参考值，m/s²
    h0 = 13.0  # mm

    # 表 1：直接测量周期，计时器误差限和卷尺误差限来自报告模板。
    t1 = values(range(4, 10), [1]).ravel()
    l1 = (572.0 - h0) / 1000
    mean1 = t1.mean()
    u_ta = t1.std(ddof=1) / math.sqrt(len(t1))
    u_tb = 0.00004 / math.sqrt(3)
    u_t = math.hypot(u_ta, u_tb)
    u_l = 0.0004 / math.sqrt(3)
    g1 = 4 * pi**2 * l1 / mean1**2
    u_g1 = g1 * math.hypot(u_l / l1, 2 * u_t / mean1)
    l_min = u_l / (0.005 / math.sqrt(2))
    t_min = u_tb / (0.005 / (2 * math.sqrt(2)))

    # 表 2：逐行由六次周期重新计算平均值，避免采用预先四舍五入的 K 列。
    h2 = values(range(4, 15), [4]).ravel()
    l2 = (h2 - h0) / 1000
    trials2 = values(range(4, 15), range(5, 11))
    mean2 = trials2.mean(axis=1)
    sq2 = mean2**2
    k2, b2, r22, _ = fit_line(l2, sq2)
    g2 = 4 * pi**2 / k2
    figure(l2, sq2, k2, b2, "摆长 L / m", "平均周期平方 T² / s²",
           "表 2：周期平方与摆长", "表2_T2-L.png",
           note=f"T² = {k2:.6f} L + {b2:.6f}\nR² = {r22:.6f}")

    # 表 3：讲义式 (12)/(20) 中根号覆盖 T/T0 - 8/9 的整体。
    h3, l3_mm, x0, t0 = 552.0, 552.0 - h0, 196.5, 1.47838
    x3 = values(range(5, 16), [13]).ravel()
    trials3 = values(range(5, 16), range(14, 20))
    mean3 = trials3.mean(axis=1)
    y3 = np.sqrt(mean3 / t0 - 8 / 9)
    k3, b3, r23, _ = fit_line(x3, y3)
    k3_theory = 3 / (16 * l3_mm)
    b3_theory = 1 / 3 - 3 * x0 / (16 * l3_mm)
    theta3 = np.degrees(np.arccos(1 - (x3 - x0) / l3_mm))
    figure(x3, y3, k3, b3, "小球参考高度 x / mm",
           r"$\sqrt{T/T_0-8/9}$", "表 3：周期与摆角",
           "表3_摆角关系.png", theory=(k3_theory, b3_theory),
           note=f"拟合斜率 = {k3:.7f} mm$^{{-1}}$\nR² = {r23:.5f}")

    # 表 4：只使用实际提供的四个摆球及已测 Ti，不推算 ti。
    ball = values(range(3, 7), [22]).ravel().astype(int)
    mass_g = values(range(3, 7), [23]).ravel()
    diameter_mm = values(range(3, 7), [24]).ravel()
    ti = values(range(3, 7), [25]).ravel()
    volume_m3 = pi / 6 * (diameter_mm / 1000) ** 3
    rho_s = (mass_g / 1000) / volume_m3
    inv_rho_s = 1 / rho_s
    sq4 = ti**2
    k4, b4, r24, cov4 = fit_line(inv_rho_s, sq4)
    rho_air = k4 / b4
    grad = np.array([1 / b4, -k4 / b4**2])
    u_rho_fit = math.sqrt(grad @ cov4 @ grad)
    figure(inv_rho_s * 1000, sq4, k4 / 1000, b4,
           r"摆球密度倒数 $1/\rho_s$ / $(10^{-3}\,\mathrm{m^3\,kg^{-1}})$", r"周期平方 $T_i^2$ / s²",
           "表 4：周期平方与摆球密度倒数", "表4_空气密度.png",
           note=f"ρ = k/b = {rho_air:.3f} kg/m³\nR² = {r24:.5f}")

    def row(items):
        return "| " + " | ".join(str(i) for i in items) + " |"

    lines = [
        r"# 单摆法测重力加速度和空气密度：数据测量与分析、思考题",
        r"",
        r"> 数据来源：`data.xlsx`；原理与表格格式依据《单摆法测重力加速度和空气密度-讲义.pdf》和《报告.pdf》。长沙重力加速度参考值取报告所给的 9.791 m/s²。所有拟合均使用未提前取整的六次测量平均值。",
        r"",
        r"## 数据测量与分析",
        r"",
        r"### 1. 误差均分设计与单点测量重力加速度（表 1）",
        r"",
        rf"梁夹距 $h_0=13$ mm，梁球距 $h=572$ mm，故摆长 $L=h-h_0={l1:.3f}$ m。计时器误差限为 0.00004 s，钢卷尺对摆长的误差限按报告给出的 0.4 mm 处理。",
        r"",
        row(["序号", *range(1, 7), "平均值"]),
        row(["---"] * 8),
        row(["周期 $T$/s", *(f"{v:.5f}" for v in t1), f"{mean1:.6f}"]),
        r"",
        r"由 $u_g/g=\sqrt{(u_L/L)^2+(2u_T/T)^2}$，按两项等额分配 0.5% 的目标相对不确定度，得 $u_L/L\leq0.5\%/\sqrt2$、$u_T/T\leq0.5\%/(2\sqrt2)$。设计阶段仅计 B 类不确定度：",
        r"",
        rf"- $u_L=0.4/\sqrt3={u_l*1000:.4f}$ mm，最小设计摆长 $L_{{\min}}={l_min*1000:.1f}$ mm。实际 $L=559$ mm，大于该下限。",
        rf"- $u_{{TB}}=0.00004/\sqrt3={u_tb:.7f}$ s，直接测一个周期时的最小设计周期 $T_{{\min}}={t_min:.5f}$ s。实际约 1.498 s，计时器分辨率满足要求；此处使用直接周期数据，故无需累积周期数 $N$ 或累计时间 $t$。",
        r"",
        r"实际测量中，周期平均值的 A 类标准不确定度按 $u_{TA}=\sqrt{\sum_i(T_i-\bar T)^2/[n(n-1)]}$ 计算，而不是对偏差直接求和：",
        r"",
        r"$$",
        rf"u_{{TA}}={u_ta:.7f}\ \mathrm{{s}},\quad u_{{TB}}={u_tb:.7f}\ \mathrm{{s}},\quad u_T=\sqrt{{u_{{TA}}^2+u_{{TB}}^2}}={u_t:.7f}\ \mathrm{{s}}",
        r"$$",
        r"",
        r"$$",
        rf"g=\frac{{4\pi^2L}}{{\bar T^2}}={g1:.4f}\ \mathrm{{m/s^2}},\qquad u_g=g\sqrt{{(u_L/L)^2+(2u_T/\bar T)^2}}={u_g1:.4f}\ \mathrm{{m/s^2}}",
        r"$$",
        r"",
        rf"结果为 **$g=({g1:.3f}\pm{u_g1:.3f})$ m/s²**（标准不确定度，$k=1$）。相对标准不确定度为 **{u_g1/g1*100:.3f}%**，满足 0.5% 的设计指标；相对长沙参考值的偏差为 **{(g1/g_ref-1)*100:+.3f}%**。若将 $h$ 与 $h_0$ 视为两次独立卷尺读数，其长度不确定度应改为 $\sqrt2\times0.4/\sqrt3$ mm，结论仍满足设计指标。",
        r"",
        r"### 2. 周期与摆长的关系（表 2）",
        r"",
        rf"$h_0={h0:g}$ mm，$L=(h-h_0)/1000$ m。每个摆长测 6 次，先求 $\bar T$，再取 $\bar T^2$。",
        r"",
        row(["$h$/mm", "$L$/m", "T₁/s", "T₂/s", "T₃/s", "T₄/s", "T₅/s", "T₆/s", "$\\bar T$/s", "$\\bar T^2$/s²"]),
        row(["---"] * 10),
    ]
    for h, l, trials, t, sq in zip(h2, l2, trials2, mean2, sq2):
        lines.append(row([f"{h:.0f}", f"{l:.3f}", *(f"{v:.5f}" for v in trials), f"{t:.6f}", f"{sq:.6f}"]))
    lines += [
        r"",
        r"![表 2：周期平方与摆长拟合](表2_T2-L.png)",
        r"",
        rf"最小二乘拟合（$L$ 用 m）：$\bar T^2=({k2:.6f}\ \mathrm{{s^2/m}})L+({b2:.6f}\ \mathrm{{s^2}})$，$R^2={r22:.6f}$。由斜率得 $g=4\pi^2/k={g2:.4f}$ m/s²，与长沙参考值的相对偏差为 **{(g2/g_ref-1)*100:+.3f}%**。截距接近零；未强制拟合线过原点，以便观察固定偏差。该组摆长 279–532 mm 均短于表 1 的 559 mm，与讲义建议“取比初始摆长更长”的步骤不一致，但不影响按实测数据进行本次拟合。",
        r"",
        r"### 3. 周期与摆角的关系（表 3）",
        r"",
        rf"$h_0=13$ mm，$h=552$ mm，$L={l3_mm:.0f}$ mm，$x_0={x0}$ mm，使用原始数据给出的 $T_0={t0:.5f}$ s。由 $\sin^2(\theta/2)=(x-x_0)/(2L)$，令 $y=\sqrt{{\bar T/T_0-8/9}}$。讲义的四次项近似给出 $y=3x/(16L)+[1/3-3x_0/(16L)]$。",
        r"",
        row(["$x$/mm", "T₁/s", "T₂/s", "T₃/s", "T₄/s", "T₅/s", "T₆/s", "$\\bar T$/s", "$y$"]),
        row(["---"] * 9),
    ]
    for x, trials, t, y in zip(x3, trials3, mean3, y3):
        lines.append(row([f"{x:.0f}", *(f"{v:.5f}" for v in trials), f"{t:.6f}", f"{y:.6f}"]))
    lines += [
        r"",
        r"![表 3：摆角关系拟合与理论线](表3_摆角关系.png)",
        r"",
        rf"实测拟合为 $y=({k3:.7f}\ \mathrm{{mm^{{-1}}}})x+{b3:.6f}$，$R^2={r23:.5f}$。理论斜率 $k_0={k3_theory:.7f}$ mm⁻¹、截距 $b_0={b3_theory:.6f}$；实测斜率比理论值低 **{(1-k3/k3_theory)*100:.1f}%**，截距高 **{(b3/b3_theory-1)*100:.1f}%**。周期随 $x$ 增大而增大，趋势与有限摆角修正一致，但斜率并未严格符合近似式。",
        rf"按 $\theta=\arccos[1-(x-x_0)/L]$ 换算，本组初始角约为 **{theta3.min():.1f}°–{theta3.max():.1f}°**。较大摆角下，四次项截断、释放状态和振幅在计时过程中的衰减都可能造成偏差，因此只能说数据支持周期随摆角增大的趋势，不能声称精确验证了理论系数。",
        r"",
        r"### 4. 测量空气密度（表 4）",
        r"",
        r"已直接测得各摆球的 $T_i$，因此按要求删去 $t_i$ 列，也不假定累计振动次数。原始文件只有 4 个摆球的数据，第 5 个空行不参与拟合。按球体模型计算 $V=\pi d^3/6$，$\rho_{si}=m_i/V_i$。",
        r"",
        row(["摆球", "$m_i$/g", "$d_i$/mm", "$V_i$/cm³", "$\\rho_{si}$/kg·m⁻³", "$1/\\rho_{si}$/m³·kg⁻¹", "$T_i$/s", "$T_i^2$/s²"]),
        row(["---"] * 8),
    ]
    for i, m, d, v, r, inv, t, sq in zip(ball, mass_g, diameter_mm, volume_m3, rho_s, inv_rho_s, ti, sq4):
        lines.append(row([i, f"{m:.2f}", f"{d:.1f}", f"{v*1e6:.6f}", f"{r:.2f}", f"{inv:.9f}", f"{t:.5f}", f"{sq:.8f}"]))
    lines += [
        r"",
        r"![表 4：周期平方与摆球密度倒数拟合](表4_空气密度.png)",
        r"",
        rf"以 $X=1/\rho_s$（m³/kg）、$Y=T_i^2$（s²）作线性拟合，得到 $Y=kX+b$，其中 $k={k4:.6f}$ s²·kg·m⁻³，$b={b4:.8f}$ s²，$R^2={r24:.5f}$。因此 **$\rho_{{air}}=k/b={rho_air:.3f}$ kg/m³**。",
        rf"仅从 4 个拟合点的残差估计，$\rho_{{air}}$ 的标准误差约为 {u_rho_fit:.3f} kg/m³；这不包含质量、直径、摆长、不同球安装位置以及 $T_i$ 的测量不确定度，不能当作完整测量不确定度。四个周期仅相差约 {(ti.max()-ti.min())*1000:.2f} ms，系统偏差足以明显改变斜率。讲义的一阶近似要求 $\rho_{{air}}/\rho_s\ll1$，本组最轻密度球的比值约 {rho_air/rho_s.min():.4f}，形式上满足该条件。",
        r"",
        r"## 思考题",
        r"",
        r"### 1. 为什么单摆周期与摆球质量无关？本实验中是否可用任意质量的球？",
        r"",
        r"理想单摆的切向运动方程是 $mL\ddot\theta=-mg\sin\theta$。两边约去 $m$，小角度下得到 $\ddot\theta+(g/L)\theta=0$，故 $T=2\pi\sqrt{L/g}$，与质量无关。实际实验不能任取质量：球太轻时空气浮力、阻力和摆线质量相对影响变大；球太重可能使摆线伸长、夹具滑动。还应使球尺寸相对摆长足够小，并保证光电门能稳定触发。本实验测空气密度正是利用不同密度球的非理想修正。",
        r"",
        r"### 2. 若摆线有质量不可忽略，周期公式应如何修正？",
        r"",
        r"将质量为 $M$ 的摆球视为距悬点 $L$ 的质点，将长度 $l$、质量 $m_l$ 的均匀摆线视为细杆。小角度下总转动惯量为 $I=ML^2+m_l l^2/3$，重力回复力矩系数为 $g(ML+m_l l/2)$，故",
        r"",
        r"$$T=2\pi\sqrt{\frac{ML^2+m_l l^2/3}{g(ML+m_l l/2)}}.$$",
        r"",
        r"当摆线延伸到球心、$l\approx L$ 时，$T\approx2\pi\sqrt{\frac{L}{g}\frac{M+m_l/3}{M+m_l/2}}$。这说明不能简单把摆线质量全部加到摆球上；摆线各部分到悬点的距离不同，应按物理摆的转动惯量和质心位置修正。",
        r"",
    ]
    (ROOT / "单摆实验_数据测量与分析_思考题.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"表1 g={g1:.6f}±{u_g1:.6f} m/s²")
    print(f"表2 g={g2:.6f} m/s², R²={r22:.6f}")
    print(f"表3 k={k3:.9f} mm^-1, R²={r23:.6f}")
    print(f"表4 ρ={rho_air:.6f} kg/m³, R²={r24:.6f}")


if __name__ == "__main__":
    main()
